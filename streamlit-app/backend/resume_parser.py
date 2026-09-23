"""
Resume file text extraction for the AI Resume Screener.

Converts an uploaded resume file (.txt, .pdf, or .docx) into plain text
that the analyzer can consume. Keeping this out of the Streamlit layer
preserves the UI/domain separation described in `app.py`.
"""

import io

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = ["txt", "pdf", "docx"]


class ResumeParseError(Exception):
    """Raised when an uploaded resume file cannot be converted to text."""


def _extract_pdf(data: bytes) -> str:
    """Extract text from every page of a PDF."""
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - pypdf raises varied types
        raise ResumeParseError(f"Could not read PDF: {exc}") from exc
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    """Extract text from a .docx file's paragraphs and tables."""
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises varied types
        raise ResumeParseError(f"Could not read Word document: {exc}") from exc

    lines = [paragraph.text for paragraph in document.paragraphs]
    # Many resume templates lay out sections in tables.
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_resume_text(filename: str, data: bytes) -> str:
    """Convert an uploaded resume file into plain text.

    Args:
        filename: The uploaded file's name, used to detect its type.
        data: The raw file bytes.

    Returns:
        The extracted text, stripped of surrounding whitespace.

    Raises:
        ResumeParseError: If the file type is unsupported, the file is
            unreadable, or it contains no extractable text (e.g. a
            scanned image-only PDF).
    """
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension == "txt":
        text = data.decode("utf-8", errors="ignore")
    elif extension == "pdf":
        text = _extract_pdf(data)
    elif extension == "docx":
        text = _extract_docx(data)
    else:
        raise ResumeParseError(
            f"Unsupported file type '{filename}'. "
            f"Upload one of: {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    text = text.strip()
    if not text:
        raise ResumeParseError(
            "No text could be extracted from this file. If it's a scanned "
            "PDF, paste the resume text instead."
        )
    return text
